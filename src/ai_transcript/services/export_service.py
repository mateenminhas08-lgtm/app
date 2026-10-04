from __future__ import annotations

import shutil
from pathlib import Path

from ai_transcript.models import TranscriptSession
from ai_transcript.utils import safe_filename


class ExportError(RuntimeError):
    pass


class ExportService:
    @staticmethod
    def export_transcript(session: TranscriptSession, destination: Path) -> Path:
        extension = destination.suffix.lower()
        source = session.transcript_md if extension == ".md" else session.transcript_txt
        if extension not in {".txt", ".md"}:
            raise ExportError("Transcript export must use .txt or .md.")
        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
            edited = session.directory / "edited-transcript.txt"
            if edited.exists():
                edited_text = edited.read_text(encoding="utf-8").rstrip()
                content = (
                    f"# {session.title}\n\n{edited_text}\n"
                    if extension == ".md"
                    else edited_text + "\n"
                )
                destination.write_text(content, encoding="utf-8")
            else:
                shutil.copy2(source, destination)
            return destination
        except OSError as exc:
            raise ExportError("The transcript could not be exported to that location.") from exc

    @staticmethod
    def export_notes(notes: str, title: str, destination: Path) -> Path:
        extension = destination.suffix.lower()
        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
            if extension == ".txt" or extension == ".md":
                destination.write_text(notes.rstrip() + "\n", encoding="utf-8")
            elif extension == ".docx":
                ExportService._write_docx(notes, title, destination)
            else:
                raise ExportError("Notes export must use .txt, .md, or .docx.")
            return destination
        except ExportError:
            raise
        except OSError as exc:
            raise ExportError("The notes could not be exported to that location.") from exc

    @staticmethod
    def _write_docx(markdown: str, title: str, destination: Path) -> None:
        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError as exc:
            raise ExportError("Word export requires the python-docx package.") from exc

        document = Document()
        document.core_properties.title = safe_filename(title)
        styles = document.styles
        styles["Normal"].font.name = "Aptos"
        styles["Normal"].font.size = Pt(11)

        for raw_line in markdown.splitlines():
            line = raw_line.rstrip()
            if not line:
                document.add_paragraph()
            elif line.startswith("### "):
                document.add_heading(line[4:], level=3)
            elif line.startswith("## "):
                document.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                document.add_heading(line[2:], level=1)
            elif line.startswith(("- ", "* ")):
                document.add_paragraph(line[2:], style="List Bullet")
            elif len(line) > 3 and line[:2].isdigit() and line[2:4] in {". ", ") "}:
                document.add_paragraph(line[4:], style="List Number")
            else:
                document.add_paragraph(line)
        document.save(destination)
